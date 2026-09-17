"""Source-level regression cases with values independently read from saved cells/HTML."""
import json
import unittest
import collections
import math
import openpyxl
from card_effects import unique_extras
from measurement_sources import (ROOT, WORKBOOK, HTML, extract_sources, extract_loopacord,
                                 extract_fujikiseki, parse_race_card_blocks, parse_race_runs, parse_focus)


class MeasurementSourceTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.cards = json.loads((ROOT / 'data/cards.json').read_text())
        cls.characters = json.loads((ROOT / 'data/characters.json').read_text())
        cls.formulas = openpyxl.load_workbook(WORKBOOK, data_only=False)
        cls.values = openpyxl.load_workbook(WORKBOOK, data_only=True)

    @classmethod
    def tearDownClass(cls):
        cls.formulas.close()
        cls.values.close()

    def test_all_measurement_blocks_are_extracted_with_explicit_eligibility(self):
        rows, templates = extract_loopacord(self.formulas, self.values, self.cards)
        self.assertEqual(len(rows), 178)
        self.assertEqual(len(templates), 3)
        self.assertEqual(sum(r['eligible'] for r in rows), 120)
        hello = {r['lb']: r for r in rows if r['card'] == 30052}
        self.assertEqual(set(hello), {'0LB', '1LB', '2LB', '3LB', 'MLB'})
        # The omitted right-hand table has W9:AA9 = 21,21,21,47,21 and AC9 = 360.
        self.assertEqual([hello['2LB'][s] for s in ['speed', 'stamina', 'power', 'guts', 'wit']], [21, 21, 21, 47, 21])
        self.assertEqual(hello['2LB']['sp'], 360)
        self.assertEqual(hello['2LB']['source_ref'], 'Grand Live Card Data!R9')
        self.assertTrue(hello['2LB']['eligible'])
        self.assertFalse(hello['0LB']['eligible'])
        self.assertIn('default-condition', hello['0LB']['excluded_reason'])
        self.assertFalse(next(r for r in rows if r['card'] == 10083)['eligible'])
        self.assertEqual(sum(r['updated'] == 'No' and r['runs'].endswith('+') for r in rows), 22)
        self.assertTrue(all(r['updated'] == 'Yes' and int(r['runs'][:-1]) >= 10 for r in rows if r['eligible']))

    def test_html_keeps_ids_and_distinguishes_medians_from_tooltip_means(self):
        rows = extract_fujikiseki(HTML.read_text(), self.cards)
        self.assertEqual(len(rows), 136)
        by_id = {r['cardId']: r for r in rows}
        self.assertEqual(by_id[30052]['sp_median'], 371)
        self.assertEqual(by_id[30052]['sp_mean'], 358.5)
        self.assertEqual(by_id[30052]['level_median'], 45)
        self.assertEqual(by_id[30052]['level_mean'], 43)
        self.assertEqual(by_id[30074]['sp_median'], 141)
        self.assertEqual(by_id[30074]['sp_mean'], 137.1)
        self.assertEqual(by_id[30074]['level_median'], 50)
        self.assertEqual(by_id[30074]['level_min'], 20)
        self.assertFalse(any(r['eligible'] for r in rows), 'median level 50 cannot establish every run was MLB')
        self.assertEqual(sum(len(r['empty_source_rows']) for r in rows), 398)

    def test_extracted_artifacts_match_original_sources(self):
        extract_sources(self.cards, check=True)

    def test_changed_card_source_shape_fails_instead_of_dropping_rows(self):
        ws = self.formulas['Grand Live Card Data']
        before = ws['R9'].value
        try:
            ws['R9'] = None
            with self.assertRaisesRegex(ValueError, 'missing source support-card ID'):
                extract_loopacord(self.formulas, self.values, self.cards)
        finally:
            ws['R9'] = before
        before, marker, after = HTML.read_text().partition('<tr data-lb=all data-card=30052 ')
        self.assertTrue(marker)
        self.assertIn('data-v=371', after)
        text = before + marker + after.replace('data-v=371', 'data-v=oops', 1)
        with self.assertRaises(ValueError):
            extract_fujikiseki(text, self.cards)

    def test_workbook_run_ids_totals_and_outfits_are_aligned(self):
        cards = parse_race_card_blocks(self.values)
        runs = parse_race_runs(self.values, cards, self.characters)
        self.assertEqual([len(runs[r]) for r in [28, 23]], [11, 11])
        self.assertEqual(runs[28][1]['event'], [692, 291, 401, 330, 455])
        self.assertEqual(runs[28][1]['card'], [417, 247, 402, 188, 159])
        self.assertEqual(runs[28][1]['total'], [1109, 538, 803, 518, 614])
        self.assertEqual(runs[28][1]['growth'], [20, 0, 10, 0, 0])
        self.assertEqual(runs[28][7]['growth'], [0, 0, 0, 10, 20])
        self.assertEqual(runs[23][1]['growth'], [0, 12, 12, 0, 6])
        self.assertEqual(runs[23][3]['traineeCardId'], 101502)
        self.assertEqual(runs[23][3]['growth'], [14, 8, 0, 0, 8])
        focus = parse_focus(self.values)
        self.assertEqual(focus['stamina'][0], round(1070 / 1083, 4))
        self.assertEqual(focus['stamina'][1], round(543 / 501, 4))
        self.assertEqual(focus['sprint'][3], round(545 / 519, 4))
        self.assertEqual(focus['sprint'][4], round(613 / 593, 4))

    def test_misaligned_runs_and_unmapped_trainees_fail(self):
        ws = self.values['Race Schedule Data']
        before = ws['N4'].value
        try:
            ws['N4'] = 99
            with self.assertRaisesRegex(ValueError, 'run IDs differ'):
                parse_race_card_blocks(self.values)
        finally:
            ws['N4'] = before
        ws = self.values['Race Schedule']
        before = ws['AD35'].value
        try:
            ws['AD35'] = 'unmapped outfit'
            with self.assertRaisesRegex(ValueError, 'unmatched trainee'):
                parse_race_runs(self.values, parse_race_card_blocks(self.values), self.characters)
        finally:
            ws['AD35'] = before

    def test_raw_stat_units_final_conversion_and_sample_variance(self):
        runs = parse_race_runs(self.values, parse_race_card_blocks(self.values), self.characters)
        first = runs[28][1]
        self.assertEqual(first['base'][0], 112)
        self.assertEqual(first['inheritance'][0], 120)
        self.assertEqual(first['final'][0], 1270)
        groups = collections.defaultdict(list)
        for race_count, rows in runs.items():
            for row in rows.values():
                groups[race_count, row['uma']].append(row['rawTotal'])
                raw = [a + b + c for a, b, c in zip(row['rawTotal'], row['base'], row['inheritance'])]
                displayed = [math.floor(v if v <= 1200 else 1200 + (v - 1200) / 2) for v in raw]
                self.assertTrue(all(abs(a - b) <= 1 for a, b in zip(displayed, row['final'])))
        repeated = [rows for rows in groups.values() if len(rows) > 1]
        self.assertEqual(sum(len(rows) - 1 for rows in repeated), 15)
        expected = []
        for stat in range(5):
            sse = 0
            for rows in repeated:
                mean = sum(row[stat] for row in rows) / len(rows)
                sse += sum((row[stat] - mean) ** 2 for row in rows)
            expected.append(math.sqrt(sse / 15))
        model = json.loads((ROOT / 'data/stat-model.json').read_text())
        for a, b in zip(expected, model['sigma']):
            self.assertAlmostEqual(a, b)

    def test_unknown_compound_payload_is_not_applied_by_the_fit(self):
        card = {'unique': {'fromLb': 1, 'effects': [{'type': 101, 'value': 80, 'value_1': 8, 'value_2': 10}]}}
        self.assertEqual(unique_extras(card, 0, .75), {})
        self.assertEqual(unique_extras(card, 1, .75), {8: 7.5})
        card['unique']['effects'][0]['description_en'] = 'extra descriptive text'
        self.assertEqual(unique_extras(card, 1, .75), {8: 7.5})
        card['unique']['effects'][0]['new_condition'] = 99
        self.assertEqual(unique_extras(card, 1, .75), {})


if __name__ == '__main__':
    unittest.main()
