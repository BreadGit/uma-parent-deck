"""Extract measurements from the vendored workbook and saved HTML without name matching.

Extraction preserves ineligible measurements. Eligibility is a separate explicit decision.
Unexpected populated structures fail instead of silently dropping source rows.
"""
import argparse
import collections
import csv
import hashlib
import io
import json
import math
from html.parser import HTMLParser
from html import unescape
from pathlib import Path
import re
import unicodedata
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / 'docs/umamusume'
WORKBOOK = SOURCE / 'loopacord-independent-training-research.xlsx'
HTML = SOURCE / 'fujikiseki-insights-2026-09-04.html'
STATS = ['speed', 'stamina', 'power', 'guts', 'wit']
LB = {'0LB': 0, '1LB': 1, '2LB': 2, '3LB': 3, 'MLB': 4}
TYPES = {'Friend': 'pal', 'Pal': 'pal', 'Group': 'group', **{s.title(): s for s in STATS}}
CARD_HEADERS = ['Card', 'Card Name', 'LB', 'Card Type', 'Score', 'Speed', 'Stamina', 'Power', 'Guts', 'Wit', 'Total Stat', 'SP', 'RB', 'Runs Recorded', 'Updated Stats']
CSV_FIELDS = ['card', 'name', 'lb', 'type', 'score', *STATS, 'total', 'sp', 'rb', 'runs', 'updated', 'source_ref', 'eligible', 'excluded_reason']


def require(condition, message):
    if not condition:
        raise ValueError(message)


def numeric(value, context):
    require(isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value), f'{context}: expected finite number, got {value!r}')
    return float(value)


def load_workbooks(path=WORKBOOK):
    return (openpyxl.load_workbook(path, read_only=True, data_only=False),
            openpyxl.load_workbook(path, read_only=True, data_only=True))


def extract_loopacord(formula_wb, value_wb, cards):
    sheet = 'Grand Live Card Data'
    formulas = list(formula_wb[sheet].iter_rows(values_only=True))
    values = list(value_wb[sheet].iter_rows(values_only=True))
    require(len(formulas) == len(values), 'workbook formula and cached-value rows differ')
    header_locations = [(i, j) for i, row in enumerate(formulas) for j, v in enumerate(row) if v == 'Card Name']
    require(header_locations, 'missing card-table headers')
    output, empty_templates, covered = [], [], set()
    lookup = {c['id']: c for c in cards}
    for header_row, name_col in header_locations:
        start = name_col - 1
        require(list(formulas[header_row][start:start + 15]) == CARD_HEADERS, f'{sheet} row {header_row + 1}: changed card-table header')
        for i in range(header_row + 1, len(formulas)):
            raw = list(formulas[i][start:start + 15])
            cached = list(values[i][start:start + 15])
            if len(raw) < 15:
                raw += [None] * (15 - len(raw))
                cached += [None] * (15 - len(cached))
            if raw[1] == 'Card Name':
                break
            image = re.search(r'support_card_s_(\d+)\.png', str(raw[0] or ''))
            if not image and all(v is None for v in raw):
                continue
            ref = f'{sheet}!{openpyxl.utils.get_column_letter(start + 1)}{i + 1}'
            # Formatted empty templates have image/score/total formulas but no measurements.
            if image and all(raw[j] is None for j in [1, 2, 3, *range(5, 10), 11, 12, 13, 14]):
                empty_templates.append(dict(sourceRef=ref, cardId=int(image.group(1)), reason='empty formatted row'))
                covered.add((i, start))
                continue
            if not image and not isinstance(raw[1], str):
                require(not any(isinstance(v, (int, float)) for v in raw[5:10] + raw[11:14]), f'{ref}: unrecognized populated measurement')
                continue
            require(image is not None, f'{ref}: missing source support-card ID')
            card_id = int(image.group(1))
            require(card_id in lookup, f'{ref}: card ID {card_id} absent from imported cards')
            require(isinstance(raw[1], str) and re.search(r' (R|SR|SSR)$', raw[1]), f'{ref}: invalid card name')
            rarity = raw[1].split()[-1]
            require(lookup[card_id]['rarity'] == rarity, f'{ref}: source ID/rarity mismatch')
            typ = ''.join(c for c in unicodedata.normalize('NFD', str(raw[3])) if not unicodedata.combining(c))
            require(typ in TYPES and lookup[card_id]['type'] == TYPES[typ], f'{ref}: source ID/type mismatch')
            require(raw[2] in LB, f'{ref}: unknown limit break {raw[2]!r}')
            require(raw[14] in ('Yes', 'No'), f'{ref}: unknown Updated Stats marker {raw[14]!r}')
            run_match = re.fullmatch(r'(\d+)(?:\.0)?(\+)?', str(raw[13]))
            require(run_match is not None, f'{ref}: unknown run-count format {raw[13]!r}')
            runs = int(run_match.group(1))
            reasons = []
            if raw[14] != 'Yes':
                reasons.append('not updated to default-condition stats')
            if runs < 10 or not run_match.group(2):
                reasons.append('not marked well-tested with at least 10 runs')
            if card_id == 10083:
                reasons.append('Light Hello R changes the reference deck condition')
            row = dict(zip(CSV_FIELDS[:15], cached))
            row.update(card=card_id, type=typ, runs=f'{runs}+' if run_match.group(2) else str(runs),
                       source_ref=ref, eligible=not reasons, excluded_reason='; '.join(reasons))
            for key in [*STATS, 'sp', 'rb', 'total']:
                row[key] = numeric(row[key], f'{ref} {key}')
            require(sum(row[s] for s in STATS) == row['total'], f'{ref}: cached total does not match stat cells')
            output.append(row)
            covered.add((i, start))
    source_images = {(i, j) for i, row in enumerate(formulas) for j, v in enumerate(row)
                     if isinstance(v, str) and 'support_card_s_' in v}
    require(source_images == covered, f'unhandled card image rows: {sorted(source_images - covered)}')
    require(len({(r['card'], r['lb']) for r in output}) == len(output), 'duplicate card/LB source measurements')
    return output, empty_templates


class CardTableParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.rows, self.row, self.cell = [], None, None

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == 'tr':
            require(self.row is None, 'unclosed source support row')
            if 'data-card' in attrs:
                require('data-lb' in attrs, 'support row has no limit-break scope')
                self.row = dict(attrs=attrs, cells=[], line=self.getpos()[0])
        elif self.row is not None and tag == 'td':
            self.cell = dict(attrs=[attrs], text=[])
            self.row['cells'].append(self.cell)
        elif self.cell is not None:
            self.cell['attrs'].append(attrs)

    def handle_data(self, text):
        if self.cell is not None:
            self.cell['text'].append(text)

    def handle_endtag(self, tag):
        if tag == 'td':
            self.cell = None
        elif tag == 'tr' and self.row is not None:
            self.rows.append(self.row)
            self.row, self.cell = None, None


def cell_value(cell, context):
    values = [float(a['data-v']) for a in cell['attrs'] if 'data-v' in a]
    require(len(values) == 1 and math.isfinite(values[0]), f'{context}: expected one numeric data-v')
    return values[0]


def extract_fujikiseki(html, cards):
    parser = CardTableParser()
    parser.feed(html)
    require(parser.row is None and parser.rows, 'incomplete source HTML card table')
    first_card_row = re.search(r'<tr\b[^>]*data-lb=', html)
    require(first_card_row is not None, 'missing support aggregate table')
    prefix = html[:first_card_row.start()]
    header = prefix[prefix.rfind('<thead'):]
    header = ' '.join(unescape(re.sub(r'<[^>]*>', ' ', header)).replace('↕', '').split())
    require(header == 'Card Runs Speed Stamina Power Guts Wit Total Skill pts Level Races', 'changed HTML support-column order')
    lookup = {c['id']: c for c in cards}
    aggregates, empty_lbs = [], collections.defaultdict(list)
    for row in parser.rows:
        attrs, cells = row['attrs'], row['cells']
        card_id = int(attrs['data-card'])
        require(attrs['data-group'] == attrs['data-card'], 'support row group/card mismatch')
        require(card_id in lookup, f'HTML card ID {card_id} absent from imported cards')
        require(attrs['data-rarity'] == lookup[card_id]['rarity'] and TYPES.get(attrs['data-type']) == lookup[card_id]['type'], f'HTML card {card_id}: ID/type mismatch')
        if attrs['data-lb'] != 'all':
            require(attrs['data-lb'] in ('0', '1', '2', '3', '4', 'detail', 'unknown'), 'unknown HTML LB scope')
            require(not cells, f'HTML card {card_id}: per-LB measurements now populated; implement explicit eligibility before importing')
            empty_lbs[card_id].append(attrs['data-lb'])
            continue
        require(len(cells) == 11, f'HTML card {card_id}: changed aggregate column count')
        text = ' '.join(''.join(cells[0]['text']).split())
        match = re.search(r'\[(.*?)\]\s+(.*?)\s+(SSR|SR|R)\s+(\w+)$', text)
        require(match is not None, f'HTML card {card_id}: unknown title layout {text!r}')
        out = dict(cardId=card_id, title=match.group(1), char=match.group(2), rarity=attrs['data-rarity'],
                   type=attrs['data-type'], runs=int(cell_value(cells[1], f'{card_id} runs')))
        for key, col in [(s, i + 2) for i, s in enumerate(STATS)] + [('sp', 8), ('level', 9), ('races', 10)]:
            cell = cells[col]
            out[f'{key}_median'] = cell_value(cell, f'{card_id} {key}')
            summaries = [re.fullmatch(r'mean ([\d.]+), range ([\d.]+) to ([\d.]+), n=(\d+)', a.get('title', '')) for a in cell['attrs']]
            summaries = [m for m in summaries if m]
            require(len(summaries) == 1, f'HTML card {card_id} {key}: changed mean/range tooltip')
            mean, low, high, count = summaries[0].groups()
            out.update({f'{key}_mean': float(mean), f'{key}_min': float(low), f'{key}_max': float(high), f'{key}_n': int(count)})
            require(float(low) <= out[f'{key}_median'] <= float(high) and int(count) <= out['runs'], f'HTML card {card_id}: invalid {key} summary')
        out['total'] = cell_value(cells[7], f'{card_id} total')
        # The median of totals need not equal the sum of individual stat medians.
        out.update(flagged=attrs.get('data-thin') == '1', source_ref=f'{HTML.name}:line {row["line"]}',
                   lb_scope='all', scenario_scope='mixed or unspecified', eligible=False,
                   excluded_reason='aggregate mixes limit breaks, levels, race counts and unspecified scenario conditions')
        aggregates.append(out)
    require(len({r['cardId'] for r in aggregates}) == len(aggregates), 'duplicate HTML aggregate card')
    require(set(empty_lbs) <= {r['cardId'] for r in aggregates}, 'per-LB row without aggregate')
    for row in aggregates:
        row['empty_source_rows'] = sorted(empty_lbs[row['cardId']])
    return aggregates


def csv_text(rows, fields):
    buf = io.StringIO(newline='')
    writer = csv.DictWriter(buf, fieldnames=fields, lineterminator='\n')
    writer.writeheader()
    writer.writerows({k: json.dumps(v, ensure_ascii=False) if isinstance(v, list) else v for k, v in row.items()} for row in rows)
    return buf.getvalue()


def extract_sources(cards, write=False, check=False):
    formulas, values = load_workbooks()
    try:
        loop, templates = extract_loopacord(formulas, values, cards)
    finally:
        formulas.close()
        values.close()
    fuji = extract_fujikiseki(HTML.read_text(), cards)
    artifacts = {
        SOURCE / 'loopacord-card-data.csv': csv_text(loop, CSV_FIELDS),
        SOURCE / 'fujikiseki-card-table.csv': csv_text(fuji, list(fuji[0])),
        SOURCE / 'fujikiseki-card-table.json': json.dumps(fuji, ensure_ascii=False, indent=1) + '\n',
    }
    for path, content in artifacts.items():
        if write:
            path.write_text(content)
        elif check:
            require(path.read_text() == content, f'{path.name} is stale; run python3 analysis/measurement_sources.py --write')
    audit = dict(loopacord=dict(measuredRows=len(loop), eligibleRows=sum(r['eligible'] for r in loop), emptyTemplates=templates),
                 fujikiseki=dict(aggregateRows=len(fuji), eligibleRows=0,
                                 emptyLbRows=sum(sum(scope in ('0', '1', '2', '3', '4') for scope in r['empty_source_rows']) for r in fuji),
                                 emptyDetailRows=sum(r['empty_source_rows'].count('detail') for r in fuji),
                                 emptyUnknownLbRows=sum(r['empty_source_rows'].count('unknown') for r in fuji)),
                 sources={p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in (WORKBOOK, HTML)})
    return loop, fuji, audit


# Source labels identify outfits, not character names. Both Biwa outfits occur.
TRAINEE_OUTFITS = {'Oguri': 100601, 'Biwa': 102301, 'XBiwa': 102302, 'Ines Fujin': 103101, 'NYOpera': 101502}
RUN_HEADER = ['Run Number', 'Speed', 'Stamina', 'Power', 'Guts', 'Wit', 'SP']


def parse_race_card_blocks(wb):
    rows = list(wb['Race Schedule Data'].iter_rows(values_only=True))
    blocks, races, current = [], None, None
    for row_index, row in enumerate(rows, 1):
        labels = [re.match(r'(23|28) Balanced G1s', value) for value in row if isinstance(value, str)]
        labels = [m for m in labels if m]
        if labels:
            require(len(labels) == 1, 'ambiguous race-count label')
            races = int(labels[0].group(1))
        starts = [j for j, value in enumerate(row) if value == 'Run Number']
        if starts:
            require(races is not None and len(starts) == 2, f'Race Schedule Data row {row_index}: expected two card tables')
            for start in starts:
                require(list(row[start:start + 7]) == RUN_HEADER, 'changed per-card race table header')
            current = dict(races=races, sourceRow=row_index, starts=starts, cards=[{}, {}])
            blocks.append(current)
        elif current:
            for slot, start in enumerate(current['starts']):
                if len(row) <= start or not isinstance(row[start], (int, float)):
                    continue
                run = int(row[start])
                require(run == row[start] and run not in current['cards'][slot], 'invalid or duplicate per-card run number')
                require(len(row[start + 1:start + 7]) == 6, 'incomplete per-card stat row')
                current['cards'][slot][run] = [numeric(v, f'Race Schedule Data row {row_index}') for v in row[start + 1:start + 7]]
    by_races = {r: [c for b in blocks if b['races'] == r for c in b['cards']] for r in (28, 23)}
    require(all(len(cards) == 6 for cards in by_races.values()), 'expected six card tables at both race counts')
    for races, cards in by_races.items():
        expected = set(cards[0])
        require(expected and expected == set(range(1, len(expected) + 1)), f'{races}: incomplete per-card run IDs')
        require(all(set(card) == expected for card in cards), f'{races}: per-card run IDs differ')
    return by_races


def parse_race_runs(wb, card_runs, characters):
    rows = list(wb['Race Schedule'].iter_rows(values_only=True))
    by_races, phases, current, races = {28: {}, 23: {}}, collections.Counter(), None, None
    for row_index, row in enumerate(rows, 1):
        cells = list(row) + [None] * 40
        labels = [re.match(r'(23|28) Balanced G1s', v) for v in cells if isinstance(v, str)]
        labels = [m for m in labels if m]
        if labels:
            races = int(labels[0].group(1))
        if cells[9] == 'Run Number':
            require(races is not None and cells[9:16] == RUN_HEADER and cells[18:25] == RUN_HEADER, 'changed race-run table header')
            phases[races] += 1
            require(phases[races] <= 3, f'{races}: unexpected extra run block')
            current = phases[races]
            if current == 3:
                require(cells[26:31] == ['Run Number', 'Race Wins', 'Crown/Tiara', 'Uma', 'Stars'], 'changed trainee metadata header')
            continue
        if not isinstance(cells[9], (int, float)):
            continue
        require(current is not None, 'race run before header')
        run = int(cells[9])
        require(cells[18] == run, f'Race Schedule row {row_index}: misaligned run IDs')
        out = by_races[races].setdefault(run, {})
        first = [numeric(v, f'Race Schedule row {row_index}') for v in cells[10:16]]
        second = [numeric(v, f'Race Schedule row {row_index}') for v in cells[19:25]]
        key = {1: 'event', 2: 'total', 3: 'final'}[current]
        require(key not in out, f'duplicate {key} run {run}')
        out[key], out[key + 'Sp'] = first[:5], first[5]
        out[key + 'SourceRef'] = f'Race Schedule!J{row_index}'
        if current == 1:
            out['card'], out['cardSp'] = second[:5], second[5]
        elif current == 3:
            require(cells[26] == run and cells[29] in TRAINEE_OUTFITS, f'Race Schedule row {row_index}: unmatched trainee metadata')
            out.update(uma=cells[29], traineeCardId=TRAINEE_OUTFITS[cells[29]], wins=int(numeric(cells[27], 'race wins')))
    lookup = {c['cardId']: c for c in characters}
    require(phases == {28: 3, 23: 3}, 'missing race-run blocks')
    for races, runs in by_races.items():
        require(set(runs) == set(card_runs[races][0]), f'{races}: deck/per-card run IDs differ')
        for run, data in runs.items():
            require(all(k in data for k in ['event', 'total', 'final', 'uma']), f'{races}/{run}: missing run block')
            require(data['traineeCardId'] in lookup, f'unknown trainee outfit {data["traineeCardId"]}')
            data['growth'] = lookup[data['traineeCardId']]['growth']
            sums = [sum(card[run][i] for card in card_runs[races]) for i in range(6)]
            require(sums == data['card'] + [data['cardSp']], f'{races}/{run}: individual cards do not sum to deck contributions')
            expected = [math.floor(v if v <= 1200 else 1200 + (v - 1200) / 2) for v in (a + b for a, b in zip(data['event'], data['card']))]
            require(expected == data['total'] and data['eventSp'] + data['cardSp'] == data['totalSp'], f'{races}/{run}: event/card/net-total mismatch')
    return by_races


def parse_focus(wb):
    rows = list(wb['Race Mode'].iter_rows(values_only=True))
    groups, current = [], None
    for row_index, row in enumerate(rows, 1):
        cells = list(row) + [None] * 14
        if cells[6] == 'Mode':
            require(cells[6:14] == ['Mode', 'Speed', 'Stamina', 'Power', 'Guts', 'Wit', 'SP', 'Total Stat Gain'], 'changed focus summary header')
            current = {}
            groups.append(current)
        elif cells[6] in ('Balanced', 'Stamina', 'Sprint'):
            require(current is not None and cells[6] not in current, 'duplicate focus summary')
            current[cells[6]] = [numeric(v, f'Race Mode row {row_index}') for v in cells[7:13]]
    require(len(groups) == 3 and all(set(g) == {'Balanced', 'Stamina', 'Sprint'} for g in groups), 'missing focus total/event/card summaries')
    total, event, card = groups
    for mode in total:
        require(all(abs(t - e - c) <= 1 for t, e, c in zip(total[mode], event[mode], card[mode])), f'{mode}: focus total does not equal event plus card summaries')
    # These are total/event/card summaries of the same runs, not independent decks.
    focus = {mode.lower(): [round(v / ref, 4) for v, ref in zip(values[:5], total['Balanced'][:5])] for mode, values in total.items()}
    return focus


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--write', action='store_true', help='regenerate vendored measurement extracts')
    args = parser.parse_args()
    _, _, audit = extract_sources(json.loads((ROOT / 'data/cards.json').read_text()), write=args.write, check=not args.write)
    print(json.dumps(audit, indent=2))
